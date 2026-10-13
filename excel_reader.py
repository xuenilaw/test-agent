"""Read one-row-per-case workbooks without editing them."""

import copy
import json
import re
from pathlib import Path

from openpyxl import load_workbook

HEADERS = ["TestCase_ID", "Target Application", "TestCase_Name", "Pre-condition",
           "Actions", "Expected", "Post action"]


def numbered_items(text: str) -> list[dict]:
    """Require explicit, consecutive numbers; don't silently discard lines."""
    parts = re.split(r"(?m)^\s*(\d+)[.)]\s+", text.strip())
    if len(parts) < 3 or parts[0].strip():
        raise ValueError("Actions and Expected must use numbered lines, starting at 1.")
    items = [{"number": int(parts[i]), "text": parts[i + 1].strip()}
             for i in range(1, len(parts), 2)]
    if [item["number"] for item in items] != list(range(1, len(items) + 1)):
        raise ValueError("Step numbers must be unique and consecutive, starting at 1.")
    if any(not item["text"] for item in items):
        raise ValueError("Numbered items cannot be empty.")
    return items


def read_case(path: Path, case_id: str) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    matches = []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            first = next(rows, ())
            headers = [str(value).strip() if value is not None else "" for value in first]
            if "TestCase_ID" not in headers:
                continue
            if any(headers.count(name) != 1 for name in HEADERS):
                raise ValueError(f"{sheet.title}: missing or duplicate required headers.")
            for row_number, row in enumerate(rows, 2):
                values = {name: row[headers.index(name)] for name in HEADERS}
                if str(values["TestCase_ID"]).strip() != case_id:
                    continue
                if any(not isinstance(v, str) or not v.strip() or v.startswith("=")
                       for v in values.values()):
                    raise ValueError(f"{sheet.title} row {row_number}: required fields must be nonempty literal text.")
                matches.append({"case_id": case_id, "name": values["TestCase_Name"].strip(),
                                "application": values["Target Application"].strip(),
                                "preconditions": values["Pre-condition"].strip(),
                                "post_action": values["Post action"].strip(),
                                "actions": numbered_items(values["Actions"]),
                                "expected": numbered_items(values["Expected"]),
                                "source": {"file": str(path.resolve()), "sheet": sheet.title,
                                           "row": row_number,
                                           "cells": {name: f"{sheet.cell(1, headers.index(name)+1).column_letter}{row_number}"
                                                     for name in HEADERS}}})
    finally:
        workbook.close()
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one {case_id} case; found {len(matches)}.")
    case = matches[0]
    if len(case["actions"]) != len(case["expected"]):
        raise ValueError("This milestone requires one numbered Expected entry per source step.")
    return case


def apply_amendments(case: dict, path: Path | None) -> tuple[dict, list[dict]]:
    effective = copy.deepcopy(case)
    applied = []
    if path is None:
        return effective, applied
    amendment = json.loads(path.read_text(encoding="utf-8")).get(case["case_id"], {})
    for column, key in (("Actions", "actions"), ("Expected", "expected")):
        for number, change in amendment.get(column, {}).items():
            index = int(number) - 1
            if index < 0 or index >= len(effective[key]):
                raise ValueError(f"Amendment refers to missing {column} {number}.")
            original = effective[key][index]["text"]
            if original != change["original"]:
                raise ValueError(f"{column} {number} has changed. Update or omit the amendments file.")
            effective[key][index]["text"] = change["replacement"]
            applied.append({"column": column, "number": int(number), "original": original,
                            "replacement": change["replacement"], "reason": amendment.get("reason", "")})
    return effective, applied
