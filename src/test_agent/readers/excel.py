"""Read one-row-per-case workbooks without editing them."""

import re
from pathlib import Path

from openpyxl import load_workbook

HEADERS = ["TestCase_ID", "Target Application", "TestCase_Name", "Pre-condition",
           "Actions", "Expected", "Post action"]


def read_case_ids(path: Path) -> list[str]:
    """Workbook order; IDs/header errors are suite input errors, not model inputs."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    ids = []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v).strip() if v is not None else "" for v in next(rows, ())]
            if "TestCase_ID" not in headers:
                continue
            if any(headers.count(name) != 1 for name in HEADERS):
                raise ValueError(f"{sheet.title}: missing or duplicate required headers")
            for row in rows:
                if all(v is None or v == "" for v in row):
                    continue
                value = row[headers.index("TestCase_ID")]
                if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]*", value.strip()):
                    raise ValueError("Invalid case ID")
                value = value.strip()
                if value in ids:
                    raise ValueError("Duplicate case ID: " + value)
                ids.append(value)
    finally:
        workbook.close()
    if not ids:
        raise ValueError("No cases in workbook")
    return ids


def numbered_items(text: str) -> list[dict]:
    """Require explicit, consecutive numbers; don't silently discard lines."""
    parts = re.split(r"(?m)^\s*(\d+)[.)]\s+", text.strip())
    if len(parts) < 3 or parts[0].strip():
        raise ValueError("Actions and Expected must use numbered lines, starting at 1.")
    items = [{"number": int(parts[i]), "text": parts[i + 1].strip()}
             for i in range(1, len(parts), 2)]
    if [item["number"] for item in items] != list(range(1, len(items) + 1)):
        raise ValueError("Step numbers must be unique and consecutive, starting at 1.")
    if any(not item["text"] for item in items):
        raise ValueError("Numbered items cannot be empty.")
    return items


def read_case(path: Path, case_id: str) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    matches = []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            first = next(rows, ())
            headers = [str(value).strip() if value is not None else "" for value in first]
            if "TestCase_ID" not in headers:
                continue
            if any(headers.count(name) != 1 for name in HEADERS):
                raise ValueError(f"{sheet.title}: missing or duplicate required headers.")
            for row_number, row in enumerate(rows, 2):
                values = {name: row[headers.index(name)] for name in HEADERS}
                if str(values["TestCase_ID"]).strip() != case_id:
                    continue
                if any(not isinstance(v, str) or not v.strip() or v.startswith("=")
                       for v in values.values()):
                    raise ValueError(f"{sheet.title} row {row_number}: required fields must be nonempty literal text.")
                matches.append({"case_id": case_id, "name": values["TestCase_Name"].strip(),
                                "application": values["Target Application"].strip(),
                                "preconditions": values["Pre-condition"].strip(),
                                "post_action": values["Post action"].strip(),
                                "actions": numbered_items(values["Actions"]),
                                "expected": numbered_items(values["Expected"]),
                                "source": {"file": str(path.resolve()), "sheet": sheet.title,
                                           "row": row_number,
                                           "cells": {name: f"{sheet.cell(1, headers.index(name)+1).column_letter}{row_number}"
                                                     for name in HEADERS}}})
    finally:
        workbook.close()
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one {case_id} case; found {len(matches)}.")
    case = matches[0]
    if len(case["actions"]) != len(case["expected"]):
        raise ValueError("This milestone requires one numbered Expected entry per source step.")
    return case
