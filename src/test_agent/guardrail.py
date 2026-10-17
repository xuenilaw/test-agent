"""Deterministic workbook preflight. No model, browser or credential access."""
import json
import re
import unicodedata
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from pydantic import Field, field_validator
from test_agent.schemas.plan import StrictModel
from test_agent.readers.excel import HEADERS

SECRET = re.compile(
    r'-----BEGIN [A-Z0-9 ]*PRIVATE KEY-----[\s\S]*?(?:-----END [A-Z0-9 ]*PRIVATE KEY-----|$)'
    r'|(?<!\w)(?:sk-ant-|sk-|ghp_|github_pat_|AKIA|xox[bp]-)[A-Za-z0-9_-]+'
    r'|\bBearer\s+[^\s<>"\x27]+'
    r'|\beyJ[A-Za-z0-9_-]*\.eyJ[A-Za-z0-9_-]*\.[A-Za-z0-9_-]+', re.I)
EMAIL = re.compile(r'(?<![\w.+-])[\w.+-]+@[\w-]+(?:\.[\w-]+)+', re.I)
URL = re.compile(r'\b(?:[a-z][a-z0-9+.-]*://|www\.)[^\s<>"\x27]+', re.I)


class GuardrailConfig(StrictModel):
    allowed_apps: dict[str, str] = Field(min_length=1)
    allowed_step_verbs: list[str] = Field(min_length=1)
    blocked_words: list[str] = Field(min_length=1)

    @field_validator('allowed_apps')
    @classmethod
    def apps(cls, apps):
        names = set()
        for name, url in apps.items():
            normalized = name.strip().casefold()
            parts = urlsplit(url)
            if not normalized or normalized in names or url != url.strip() or parts.scheme not in {'http','https'} or not parts.hostname or parts.username or parts.password:
                raise ValueError('Invalid or duplicate app')
            # Also validate malformed ports.
            parts.port
            names.add(normalized)
        return apps

    @field_validator('allowed_step_verbs', 'blocked_words')
    @classmethod
    def phrases(cls, values, info):
        if any(not value.strip() or value != value.strip() for value in values):
            raise ValueError('Empty or padded phrase')
        if info.field_name=='allowed_step_verbs' and any(not re.fullmatch(r'[a-zA-Z]+(?: [a-zA-Z]+)*',v) for v in values):
            raise ValueError('Invalid verb phrase')
        return values


def masked(text):
    text = SECRET.sub(lambda m:m[0][:4]+'…',str(text))
    text = EMAIL.sub(lambda m:m[0][:4]+'…',text)
    return ' '.join(text.split())[:80]


def phrase_pattern(phrase):
    return re.compile(r'(?<!\w)'+r'\s+'.join(re.escape(p) for p in phrase.split())+r'(?:s|es)?(?!\w)',re.I)


def step_text(line):
    line = line.strip()
    # Strip quotes/bullets before and after numbering; do not strip letters/digits.
    def punctuation(value):
        return value.lstrip(''.join(c for c in set(value) if unicodedata.category(c).startswith(('P','S')) or c.isspace()))
    line = punctuation(line)
    match = re.match(r'^(\d+)[.)]\s*',line)
    return (int(match[1]) if match else None), punctuation(line[match.end():] if match else line)


def check_workbook(workbook, *, url=None, case_id=None):
    """Scan every text cell (including extra columns/sheets); fail closed on inputs."""
    workbook = Path(workbook)
    result = dict(status='PASSED',violations=[],application_urls={})
    def violation(case, column, step, rule, text, message):
        result['violations'].append(dict(case=masked(case),column=masked(column),step=step,
            rule=rule,snippet=masked(text),message=message))
    try:
        config = GuardrailConfig.model_validate_json((workbook.parent/'guardrail.json').read_text())
    except Exception:
        violation('workbook','config',None,'CONFIG_ERROR','', 'Missing or invalid guardrail.json; a valid human-owned config is required.')
        result['status']='BLOCKED'; return result
    apps = {name.strip().casefold():value for name,value in config.allowed_apps.items()}
    blocked = [phrase_pattern(word) for word in config.blocked_words]
    verbs = [re.compile(r'^'+r'\s+'.join(re.escape(p) for p in verb.split())+r'(?!\w)',re.I) for verb in config.allowed_step_verbs]
    book = None
    try:
        book = load_workbook(workbook,read_only=True,data_only=False)
        ids = set()
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v).strip() if v is not None else '' for v in next(rows,())]
            case_sheet = 'TestCase_ID' in headers
            if case_sheet and any(headers.count(h)!=1 for h in HEADERS):
                raise ValueError('Headers')
            # Header/notes text must pass the same content rules as case cells.
            values = [(1,headers)] + list(enumerate(rows,2))
            for row_number,row in values:
                if all(v is None or v=='' for v in row):
                    continue
                case, app_url = 'workbook', None
                if case_sheet and row_number>1:
                    case = row[headers.index('TestCase_ID')]
                    if not isinstance(case,str) or not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]*',case.strip()) or case.strip() in ids:
                        raise ValueError('Case ID')
                    case = case.strip(); ids.add(case)
                    app = row[headers.index('Target Application')]
                    app_url = apps.get(str(app).strip().casefold())
                    if app_url is None:
                        violation(case,'Target Application',None,'ALLOWED_APP',str(app),'Target application is not allowed.')
                    else:
                        result['application_urls'][case]=app_url
                    if url is not None and (case_id is None or case==case_id) and url!=app_url:
                        violation(case,'Target Application',None,'URL_OVERRIDE',url,'URL override must equal this case\'s configured app URL.')
                for index,value in enumerate(row):
                    if not isinstance(value,str):
                        continue
                    column = headers[index] if index<len(headers) and headers[index] else 'column '+str(index+1)
                    secret_matches = list(SECRET.finditer(value))
                    def emit(rule, match, message):
                        prefix = value[:match.start()]
                        line = value[prefix.rfind('\n')+1:].split('\n',1)[0]
                        step, _ = step_text(line)
                        # A full-cell mask removes private-key bodies across line boundaries.
                        snippet = value if any('\n' in m[0] for m in secret_matches) else line
                        violation(case,column,step,rule,match[0] if rule=='SECRET' else snippet,message)
                    for match in secret_matches:
                        emit('SECRET',match,'Secret-looking string is forbidden.')
                    for match in EMAIL.finditer(value):
                        emit('EMAIL',match,'Email addresses are forbidden in Excel.')
                    for pattern in blocked:
                        for match in pattern.finditer(value):
                            emit('BLOCKED_WORD',match,'Contains blocked phrase: '+masked(match[0])+'.')
                    for match in URL.finditer(value):
                        candidate = match[0].rstrip('.,;:!?)]}')
                        try:
                            host = urlsplit('https://'+candidate if candidate.lower().startswith('www.') else candidate).hostname
                        except ValueError:
                            host = None
                        if not app_url or not host or host.casefold()!=urlsplit(app_url).hostname.casefold():
                            emit('URL_HOST',match,'URL host must exactly match the case\'s app host.')
                    for line in value.splitlines():
                        step, text = step_text(line)
                        if case_sheet and row_number>1 and column in {'Actions','Post action'} and line.strip():
                            if not any(pattern.match(text) for pattern in verbs):
                                snippet = value if any('\n' in m[0] for m in secret_matches) else line
                                violation(case,column,step,'STEP_VERB',snippet,'Step does not start with an allowed verb.')
        if not ids:
            raise ValueError('No cases')
        from test_agent.readers.excel import read_case
        for case_id in ids:
            read_case(workbook,case_id)  # Validate all required fields, even cases not selected for translation.
    except Exception:
        violation('workbook','input',None,'INPUT_ERROR','', 'Workbook is unreadable or has invalid case IDs/headers.')
    finally:
        if book: book.close()
    result['status']='BLOCKED' if result['violations'] else 'PASSED'
    if result['status']=='BLOCKED':
        result['application_urls'] = {}
    return result


def save_check(result, output):
    output = Path(output); output.mkdir(parents=True,exist_ok=True)
    (output/'guardrail.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    if result['status']=='BLOCKED':
        summary = dict(status='BLOCKED_BY_GUARDRAIL',cases=[],output_directory=str(output.resolve()),
            reason='Blocked by guardrail; nothing was executed',totals={})
        (output/'summary.json').write_text(json.dumps(summary,indent=2))
        return summary


def print_check(result):
    if result['status']=='PASSED':
        print('Guardrail PASSED — '+str(len(result['application_urls']))+' cases checked')
        return
    print('⛔ BLOCKED BY GUARDRAIL — nothing was executed')
    groups = {}
    for item in result['violations']:
        groups.setdefault((item['case'],item['step']),[]).append(item)
    for (case,number),items in groups.items():
        step = ' step '+str(number) if number is not None else ''
        columns = ', '.join(dict.fromkeys(item['column'] for item in items))
        sample = next((item for item in items if item['rule']=='STEP_VERB'),items[0])
        reasons = {}
        for item in items:
            detail = (item['message'].removeprefix('Contains blocked phrase: ').removesuffix('.')
                      if item['rule']=='BLOCKED_WORD' else item['message'])
            reasons.setdefault(item['rule'],[])
            if detail not in reasons[item['rule']]:
                reasons[item['rule']].append(detail)
        if 'STEP_VERB' in reasons:
            reasons = {'STEP_VERB':reasons['STEP_VERB'],**reasons}
        explanation = '; '.join(rule+': '+', '.join(details) for rule,details in reasons.items())
        print(f"{case}{step} [{columns}]: {sample['snippet']} → {explanation}")
    print('Fix the Excel (or, if legitimate, its adjacent guardrail.json), then run again.')
